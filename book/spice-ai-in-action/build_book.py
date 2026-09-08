"""Build the editable Word manuscript from the chapter sources."""
from pathlib import Path
import re,subprocess,json
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
root=Path(__file__).resolve().parent
INK='2E2F36';ORANGE='FF622A';MUTED='666674';PALE='F3F3F6'

def element(tag,**attrs):
 e=OxmlElement('w:'+tag)
 for k,v in attrs.items():e.set(qn('w:'+k),str(v))
 return e

def style(st,font,size,bold=False,color=INK):
 st.font.name=font;st.font.size=Pt(size);st.font.bold=bold;st.font.color.rgb=RGBColor.from_string(color)
 st.paragraph_format.widow_control=True
 rf=st.element.find('.//'+qn('w:rFonts'))
 if rf is not None:
  for attr in ['asciiTheme','hAnsiTheme','eastAsiaTheme','cstheme']:
   rf.attrib.pop(qn('w:'+attr),None)

ref=Document()
sec=ref.sections[0]
sec.page_width=Inches(7.5);sec.page_height=Inches(9.25)
sec.top_margin=Inches(.67);sec.bottom_margin=Inches(.67)
sec.left_margin=Inches(.7);sec.right_margin=Inches(.7)
sec.header_distance=Inches(.27);sec.footer_distance=Inches(.3)
style(ref.styles['Normal'],'Georgia',11)
ref.styles['Normal'].paragraph_format.line_spacing=1.15
ref.styles['Normal'].paragraph_format.space_after=Pt(7)
for name in ['Body Text','First Paragraph']:
 if name not in ref.styles:ref.styles.add_style(name,1)
 style(ref.styles[name],'Georgia',11)
 ref.styles[name].paragraph_format.space_after=Pt(7)
 ref.styles[name].paragraph_format.line_spacing=1.15
for level,size in [(1,23),(2,14),(3,11.5),(4,10.5)]:
 st=ref.styles[f'Heading {level}'];style(st,'Inter',size,True,ORANGE if level==1 else INK)
 st.paragraph_format.space_before=Pt(18 if level>1 else 5)
 st.paragraph_format.space_after=Pt(9)
 st.paragraph_format.keep_with_next=True
 if level==1:st.paragraph_format.page_break_before=True
for name,font,size in [('Source Code','Consolas',8.5),('Caption','Inter',9),('Table','Inter',9),('TOC Heading','Inter',23),('TOC 1','Inter',10)]:
 if name not in ref.styles:ref.styles.add_style(name,1)
 style(ref.styles[name],font,size,name=='TOC Heading')
ref.styles['Caption'].paragraph_format.space_after=Pt(10)
ref.styles['Caption'].font.italic=True
ref.styles['Source Code'].paragraph_format.line_spacing=1.06
ref.styles['Source Code'].paragraph_format.space_before=Pt(6)
ref.styles['Source Code'].paragraph_format.space_after=Pt(9)
ref.styles['Source Code'].paragraph_format.keep_together=False
ref.styles['Source Code'].paragraph_format.left_indent=Inches(.10)
ref.styles['Source Code'].paragraph_format.right_indent=Inches(.08)
ref.styles['TOC 1'].paragraph_format.space_after=Pt(5)
ref.styles['TOC 1'].paragraph_format.line_spacing=1.08
ref.save(root/'reference.docx')
files=list((root/'chapters').glob('*.md'))
def order(p):
 title=p.read_text().splitlines()[0]
 m=re.match(r'# (\d+)\.',title)
 if m:return(1,int(m[1]))
 if 'Preface' in title:return(0,0)
 m=re.match(r'# Appendix ([A-Z])\.',title)
 if m:return(2,ord(m[1]))
 return(3,0)
chunks=[]
for p in sorted(files,key=order):
 s=p.read_text()
 s=re.sub(r'^(# (\d+)\..+)$',lambda m:m[1]+' {#chapter-'+m[2]+'}',s,count=1,flags=re.M)
 s=re.sub(r'(!\[[^\n]+\]\(figures/[^)]+\))',r'\1{width=6.1in}',s)
 chunks.append(s)
(root/'manuscript.md').write_text('\n\n'.join(chunks))
subprocess.run(['pandoc','manuscript.md','--from=markdown','--to=docx','--standalone','--toc','--toc-depth=1','--metadata=toc-title:Contents','--syntax-highlighting=none','--reference-doc=reference.docx','--output=body.docx'],cwd=root,check=True)
doc=Document(root/'body.docx')
sec=doc.sections[0];sec.different_first_page_header_footer=True
# Place a native, editable cover before Pandoc's contents control.
position=0
for text,size,bold,color,before,after in [
 ('SPICE.AI IN ACTION',10,True,MUTED,32,72),
 ('Spice.ai\nin Action',43,True,INK,0,23),
 ('A practical guide to SQL, acceleration,\nsearch, and data-grounded AI',17,False,INK,0,25),
 ('SQL • DATA • SEARCH • AI',10,True,ORANGE,0,112),
 ('FIRST EDITION MANUSCRIPT',10,True,MUTED,0,5),
 ('September 2026',11,False,MUTED,0,0)]:
 p=doc.add_paragraph();r=p.add_run(text);r.font.name='Inter';r.font.size=Pt(size);r.bold=bold;r.font.color.rgb=RGBColor.from_string(color)
 p.paragraph_format.space_before=Pt(before);p.paragraph_format.space_after=Pt(after);p.paragraph_format.line_spacing=1.04
 p.paragraph_format.keep_together=True
 doc._element.body.remove(p._p);doc._element.body.insert(position,p._p);position+=1
p=doc.add_paragraph();p.add_run().add_break(__import__('docx').enum.text.WD_BREAK.PAGE)
doc._element.body.remove(p._p);doc._element.body.insert(position,p._p)
# Local copy-fitting keeps chapter-end references and exercises with their chapter.
compact_chapters={'2','9','16','25','26','30'}
compact=False
tight=False
for p in doc.paragraphs:
 if p.style.name=='Heading 1':
  m=re.match(r'(\d+)\.',p.text)
  compact=bool(m and m[1] in compact_chapters) or p.text.startswith('Appendix A.')
  tight=bool(m and m[1] in {'2','16','30'}) or p.text.startswith('Appendix A.')
 elif compact:
  if p.style.name in ['Body Text','First Paragraph','Normal']:
   p.paragraph_format.space_after=Pt(3.5 if tight else 5)
   p.paragraph_format.line_spacing=1.10 if tight else 1.13
  elif p.style.name=='Heading 2':
   p.paragraph_format.space_before=Pt(14 if tight else 16)
   p.paragraph_format.space_after=Pt(7 if tight else 8)

# Running furniture.
h=sec.header.paragraphs[0];h.text='SPICE.AI IN ACTION';h.alignment=WD_ALIGN_PARAGRAPH.LEFT
for r in h.runs:r.font.name='Inter';r.font.size=Pt(8);r.font.color.rgb=RGBColor.from_string(MUTED)
f=sec.footer.paragraphs[0];f.alignment=WD_ALIGN_PARAGRAPH.CENTER
r=f.add_run();r.font.name='Inter';r.font.size=Pt(8)
r._r.append(element('fldChar',fldCharType='begin'))
i=element('instrText');i.text=' PAGE ';r._r.append(i);r._r.append(element('fldChar',fldCharType='end'))
# Code shading and predictable inline code typography.
for p in doc.paragraphs:
 if p.style.name=='Source Code':
  p._p.get_or_add_pPr().append(element('shd',fill=PALE,val='clear'))
  p.paragraph_format.keep_together=len(p.text.splitlines())<=22
  for r in p.runs:r.font.name='Consolas';r.font.size=Pt(8.5)
 if p.style.name in ['Caption','Image Caption']:
  p.paragraph_format.keep_together=True
 if p._p.xpath('.//w:drawing'):
  p.paragraph_format.keep_with_next=True
 for r in p.runs:
  if r.style and r.style.name=='Verbatim Char':r.font.name='Consolas';r.font.size=Pt(9.3)
for previous,current in zip(doc.paragraphs,doc.paragraphs[1:]):
 if current.style.name=='Source Code' and previous.text.rstrip().endswith(':'):
  previous.paragraph_format.keep_with_next=True
# Tables: fixed page width, repeated headers, and shaded alternate rows.
for t in doc.tables:
 t.autofit=False
 n=len(t.columns)
 weights=[1/n]*n
 if n==2:weights=[.51,.49]
 if n==3:weights=[.22,.40,.38]
 if len(t.rows)>30 and n==3:weights=[.16,.55,.29]
 for col,w in zip(t.columns,weights):col.width=Inches(6.1*w)
 is_index = n==2 and (len(t.rows)>50 or t.cell(0,0).text=='File or directory')
 for ri,row in enumerate(t.rows):
  if ri==0:row._tr.get_or_add_trPr().append(element('tblHeader'))
  row._tr.get_or_add_trPr().append(element('cantSplit'))
  for ci,cell in enumerate(row.cells):
   cell.width=Inches(6.1*weights[ci])
   tcpr=cell._tc.get_or_add_tcPr()
   tcpr.append(element('shd',fill=INK if ri==0 else (PALE if ri%2 else 'FFFFFF'),val='clear'))
   margins=element('tcMar')
   for side in ['top','bottom','left','right']:margins.append(element(side,w=('35' if is_index else '65'),type='dxa'))
   tcpr.append(margins)
   for p in cell.paragraphs:
    p.paragraph_format.space_after=Pt(1 if is_index else 3);p.paragraph_format.space_before=Pt(1 if is_index else 2);p.paragraph_format.line_spacing=1.1
    for r in p.runs:r.font.name='Inter';r.font.size=Pt(8.5);r.bold=ri==0;r.font.color.rgb=RGBColor.from_string('FFFFFF' if ri==0 else INK)
    # Hyperlinks contain runs not exposed by python-docx's paragraph.runs.
    for run in p._p.xpath('.//w:hyperlink/w:r'):
     rp=run.find(qn('w:rPr'))
     if rp is None:rp=element('rPr');run.insert(0,rp)
     rp.append(element('rFonts',ascii='Inter',hAnsi='Inter'));rp.append(element('sz',val='17'))
     rp.append(element('color',val='FFFFFF' if ri==0 else 'BC3201'))
settings=doc.settings.element
settings.append(element('updateFields',val='true'))
props=doc.core_properties;props.title='Spice.ai in Action';props.subject='SQL, acceleration, search, and data-grounded AI';props.author='';props.last_modified_by='';props.keywords='Spice.ai, SQL, DataFusion, Arrow, Cayenne, CDC, RAG';props.comments='Original technical manuscript. Version and execution scope are documented in the preface and Appendix A.'
out=root/'Spice.ai in Action.docx';doc.save(out)
print(json.dumps({'docx':str(out),'markdown_words':len((root/'manuscript.md').read_text().split()),'chapters':36,'tables':len(doc.tables),'figures':len(doc.inline_shapes),'paragraphs':len(doc.paragraphs)},indent=2))
